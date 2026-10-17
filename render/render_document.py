"""Renders a governed Markdown document to Word (.docx) and PDF with the print template.

Template: A4, 2 cm margins, Arial 12 for every run, black text, bold headings (Word heading styles, so the
navigation pane and the PDF outline list them), thin black table borders that never pass the usable width,
and a footer with the source path and the page number. The .docx is the master; LibreOffice converts it
to PDF, so both formats always match.

Markdown it reads: front matter; # / ## / ### headings; paragraphs with **bold** and `code`; "- " bullets;
"1. " numbered steps (numbering restarts with each list); pipe tables; fenced code blocks (lines kept as
written, no Markdown inside, in a thin box); notices as a blockquote ("> **Note:** ...", indented with a
left rule); and images on their own line ("![Caption](relative/path.png)"), resolved against the source
file, scaled to the usable width (at most 18 cm high), centered and captioned "Figure N." or "Figura N."
after the lang of the front matter. A missing image stops the render.

The "> **Question this document answers:**" line stays a plain paragraph, without the quote marker.
Documents in a delivery/ folder open with a title block (the H1 as the title, the author and the
last-reviewed date), and their tables leave room for bold headers and wide letters. Other documents,
such as the ADR, keep the layout they always had.

Usage (from the workspace root):
  python orderflow-engine-tools/render/render_document.py <source.md> <output-dir> [--pdf] [--label <path>]
"""
import argparse
import hashlib
import io
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.exceptions import UnrecognizedImageError
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

FONT = 'Arial'
SIZE = Pt(12)
BLACK = RGBColor(0, 0, 0)
SOFFICE_CANDIDATES = [
    'soffice',
    r'C:\Program Files\LibreOffice\program\soffice.exe',
    '/Applications/LibreOffice.app/Contents/MacOS/soffice',
]


def front_matter(text):
    match = re.match(r'^---\n(.*?)\n---\n', text, re.S)
    data = {}
    if match:
        for line in match.group(1).split('\n'):
            key, _, value = line.partition(':')
            data[key.strip()] = value.strip()
        text = text[match.end():]
    return data, text


def style_run(run, bold=False, mono=False):
    run.font.name = FONT
    run.font.size = SIZE
    run.font.bold = bold
    run.font.color.rgb = BLACK
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if mono:
        run.font.highlight_color = None


def add_inline(paragraph, text):
    """Supports **bold** and `code` (rendered in Arial, as the template allows one font only)."""
    for token in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not token:
            continue
        if token.startswith('**'):
            style_run(paragraph.add_run(token[2:-2]), bold=True)
        elif token.startswith('`'):
            style_run(paragraph.add_run(token[1:-1]), mono=True)
        else:
            style_run(paragraph.add_run(token))


def set_cell_borders(table):
    tbl = table._tbl
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), '000000')
        borders.append(element)
    tbl.tblPr.append(borders)


# Children that follow w:pBdr inside w:pPr; Word rejects a paragraph whose properties are out of order.
PBDR_SUCCESSORS = (
    'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
    'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
    'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
    'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
    'w:rPr', 'w:sectPr', 'w:pPrChange',
)


def set_paragraph_borders(paragraph, edges, size, space):
    """Thin black paragraph borders; size in eighths of a point, space (gap to the text) in points."""
    borders = OxmlElement('w:pBdr')
    for edge in edges:
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), str(size))
        element.set(qn('w:space'), str(space))
        element.set(qn('w:color'), '000000')
        borders.append(element)
    paragraph._p.get_or_add_pPr().insert_element_before(borders, *PBDR_SUCCESSORS)


FOOTER_LABELS = {'en': ('Source', 'p.'), 'es': ('Fuente', 'pág.')}
FIGURE_LABELS = {'en': 'Figure', 'es': 'Figura'}
QUESTION_LABELS = ('**Question this document answers:**', '**Pregunta que responde este documento:**')
CHAR_EMU = 85000      # average Arial 12 character width
PADDING_EMU = 180000  # left and right cell padding
DELIVERY_WORD_SCALE = (1.3, 1.15)  # delivery tables: bold headers and wide letters never break a word
MAX_FIGURE_HEIGHT = Cm(18)
LONG_WORD = 20  # characters; a longer token (a URL, a path, an environment key) wraps inside its cell
LIST_INDENT = Twips(360)  # the same hanging indent the bullet list uses
BLOCK_INDENT = Cm(0.5)    # code blocks and notices

IMAGE = re.compile(r'^!\[([^\]]*)\]\(([^)\s]+)\)\s*$')
NUMBERED = re.compile(r'^\s*(\d+)\. (.*)$')
# A line that starts a block of its own, so it never joins the paragraph above it.
BLOCK_START = re.compile(r'^(#|\||\s*[-*] |\s*\d+\. |\s*```|>|!\[)')


def column_widths(rows, usable, header_scale=1.0, body_scale=1.0):
    """No column is narrower than its longest word, up to LONG_WORD characters, so ordinary words never break;
    the rest follows content length. The table never passes the usable width: when the minimums alone would, they
    shrink in proportion.

    The scales widen the estimate of a word in the bold header row and in the other rows (wide letters).
    """
    cols = len(rows[0])
    plain = [[re.sub(r'[*`]', '', row[c]) for c in range(cols)] for row in rows]
    minimum = [int(max(min(len(w), LONG_WORD) * (header_scale if r == 0 else body_scale)
                       for r, row in enumerate(plain) for w in (row[c].split() or [''])) * CHAR_EMU) + PADDING_EMU
               for c in range(cols)]
    if sum(minimum) > usable:
        return [int(m * usable / sum(minimum)) for m in minimum]
    content = [max(len(row[c]) for row in plain) for c in range(cols)]
    spare = max(0, usable - sum(minimum))
    extra = [max(0, content[c] - max(min(len(w), LONG_WORD) for row in plain for w in (row[c].split() or [''])))
             for c in range(cols)]
    total = sum(extra) or 1
    return [int(minimum[c] + spare * extra[c] / total) for c in range(cols)]


def add_page_number(paragraph):
    for instruction in ('PAGE',):
        run = paragraph.add_run()
        style_run(run)
        begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
        code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = instruction
        end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
        run._r.append(begin); run._r.append(code); run._r.append(end)


def add_title_block(doc, title, author, date):
    """The H1 as the title, then the author and the last-reviewed date, centered."""
    lines = [(text, bold) for text, bold in ((title, True), (author, False), (date, False)) if text]
    for n, (text, bold) in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(12 if n == len(lines) - 1 else 2)
        paragraph.paragraph_format.keep_with_next = True
        style_run(paragraph.add_run(text), bold=bold)


def add_numbered_item(doc, number, text):
    """Typed numbers with a hanging indent, so every list restarts and renders alike in Word and LibreOffice."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = LIST_INDENT
    paragraph.paragraph_format.first_line_indent = -LIST_INDENT
    paragraph.paragraph_format.tab_stops.add_tab_stop(LIST_INDENT)
    paragraph.paragraph_format.space_after = Pt(2)
    style_run(paragraph.add_run(f'{number}.\t'))
    add_inline(paragraph, text)


def add_code_block(doc, code_lines):
    """One boxed, indented paragraph; each source line on its own line, in Arial, with no Markdown inside."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = BLOCK_INDENT
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    set_paragraph_borders(paragraph, ('top', 'left', 'bottom', 'right'), size=4, space=4)
    for n, code_line in enumerate(code_lines):
        run = paragraph.add_run()
        style_run(run)
        if n:
            run.add_break()
        run.add_text(code_line.replace('\t', '    '))


def add_notice(doc, text):
    """A blockquote notice ("> **Note:** ..."): indented, with a left rule."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = BLOCK_INDENT
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    set_paragraph_borders(paragraph, ('left',), size=12, space=8)
    add_inline(paragraph, text)


def add_figure(doc, image_path, image, caption, number, lang, usable):
    """The image at the usable width (lower when it would pass the height limit), then its numbered caption."""
    width = min(int(usable), int(MAX_FIGURE_HEIGHT * image.width / image.height))
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(6)
    picture.paragraph_format.space_after = Pt(2)
    picture.paragraph_format.keep_with_next = True
    shape = picture.add_run().add_picture(str(image_path), width=width)
    # python-docx writes no wrap distances, and LibreOffice then adds its default 0.32 cm, which pushes the image
    # off the text column; zero distances keep it exactly on the column.
    for edge in ('distT', 'distB', 'distL', 'distR'):
        shape._inline.set(edge, '0')
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(8)
    style_run(label.add_run(f'{FIGURE_LABELS.get(lang, FIGURE_LABELS["en"])} {number}.'), bold=True)
    if caption:
        style_run(label.add_run(' '))
        add_inline(label, caption)


def render_docx(source, out_path, source_label, delivery=None):
    """Renders the source to out_path and returns (front matter, embedded image paths in order of use).

    delivery (by default: the source sits in a delivery/ folder) adds the title block and measures table words
    with a margin. Other documents, such as the ADR, keep the layout they always had.
    """
    source = Path(source)
    raw = io.open(source, encoding='utf-8').read().replace('\r\n', '\n')
    meta, body = front_matter(raw)
    line_offset = raw[:len(raw) - len(body)].count('\n')
    if delivery is None:
        delivery = source.parent.name == 'delivery'
    word_scale = DELIVERY_WORD_SCALE if delivery else (1.0, 1.0)
    lang = meta.get('lang', 'en')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(section, side, Cm(2.0))
    usable = section.page_width - section.left_margin - section.right_margin
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    labels = FOOTER_LABELS.get(lang, FOOTER_LABELS['en'])
    add_inline(footer, f'{labels[0]}: {source_label} · {labels[1]} ')
    add_page_number(footer)

    lines = body.split('\n')
    images = []
    figures = 0
    list_number = None  # the last number of the numbered list in progress, None outside a list
    title_pending = delivery
    i = 0
    after_table = False
    after_code = False  # the last block was a code block
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        # Indented lines belong to the list item above them; anything else ends a numbered list.
        numbered = NUMBERED.match(line)
        if not numbered and not line[:1].isspace():
            list_number = None
        was_code, after_code = after_code, False
        if line.lstrip().startswith('```'):
            after_code = was_code
            start = i
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].lstrip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            if i == len(lines):
                raise SystemExit(f'{source_label}: line {start + line_offset + 1}: code block is never closed')
            if after_code:
                # Word and LibreOffice merge adjacent paragraphs with the same borders into one box; an empty
                # paragraph between two code blocks keeps them apart.
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
            add_code_block(doc, code_lines)
            after_code = True
            after_table = False
            i += 1
            continue
        image = IMAGE.match(line)
        if image:
            caption, reference = image.group(1).strip(), image.group(2)
            image_path = (source.parent / unquote(reference)).resolve()
            where = f'{source_label}: line {i + line_offset + 1}: image {reference}'
            if not image_path.is_file():
                raise SystemExit(f'{where} not found (looked for {image_path})')
            try:
                picture = Image.from_file(str(image_path))
            except UnrecognizedImageError:
                raise SystemExit(f'{where} is not a supported image (use PNG or JPEG)')
            figures += 1
            add_figure(doc, image_path, picture, caption, figures, lang, usable)
            if image_path not in images:
                images.append(image_path)
            after_table = False
            i += 1
            continue
        if line.startswith('>'):
            quoted = []
            while i < len(lines) and lines[i].startswith('>'):
                quoted.append(re.sub(r'^>\s?', '', lines[i]).strip())
                i += 1
            for text in (' '.join(part.split('\n')) for part in '\n'.join(quoted).split('\n\n')):
                if not text.strip():
                    continue
                if text.startswith(QUESTION_LABELS):
                    paragraph = doc.add_paragraph()
                    if after_table:
                        paragraph.paragraph_format.space_before = Pt(6)
                    add_inline(paragraph, text)
                else:
                    add_notice(doc, text)
                after_table = False
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[-| :]+\|$', lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            set_cell_borders(table)
            # Column widths follow content length, so short columns do not force long ones to wrap.
            widths = column_widths(rows, usable, *word_scale)
            for c, width in enumerate(widths):
                table.columns[c].width = width
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.width = widths[c]
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    if r == 0:
                        style_run(paragraph.add_run(value.replace('**', '')), bold=True)
                    else:
                        add_inline(paragraph, value)
            after_table = True
            continue
        heading = re.match(r'^(#{1,3}) (.*)$', line)
        if heading:
            if title_pending and heading.group(1) == '#':
                add_title_block(doc, heading.group(2), license_holder(), meta.get('last-reviewed', ''))
                title_pending = False
                after_table = False
                i += 1
                continue
            # A heading style gives the Word navigation pane and the PDF outline their entries.
            paragraph = doc.add_paragraph(style=f'Heading {len(heading.group(1))}')
            paragraph.paragraph_format.space_before = Pt(8 if heading.group(1) != '#' else 0)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_with_next = True
            style_run(paragraph.add_run(heading.group(2)), bold=True)
            after_table = False
            i += 1
            continue
        if numbered:
            list_number = int(numbered.group(1)) if list_number is None else list_number + 1
            add_numbered_item(doc, list_number, numbered.group(2))
            after_table = False
            i += 1
            continue
        bullet = re.match(r'^\s*[-*] (.*)$', line)
        if bullet:
            paragraph = doc.add_paragraph(style='List Bullet')
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, bullet.group(1))
            after_table = False
            i += 1
            continue
        paragraph = doc.add_paragraph()
        if after_table:
            paragraph.paragraph_format.space_before = Pt(6)
        after_table = False
        text = line
        while i + 1 < len(lines) and lines[i + 1].strip() and not BLOCK_START.match(lines[i + 1]):
            i += 1
            text += ' ' + lines[i].strip()
        add_inline(paragraph, text)
        i += 1

    for style_name in ('List Bullet',):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = SIZE
        style.font.color.rgb = BLACK
    for level in (1, 2, 3):
        # The template's headings use theme fonts in blue; the print template allows Arial 12, black, bold only.
        style = doc.styles[f'Heading {level}']
        style.element.get_or_add_rPr().get_or_add_rFonts().attrib.clear()
        style.font.name = FONT
        style.font.size = SIZE
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = BLACK
    clean_metadata(doc, meta)
    doc.save(out_path)
    strip_application(out_path)
    return meta, images


def to_pdf(docx_path, out_dir):
    for candidate in SOFFICE_CANDIDATES:
        exe = shutil.which(candidate) or (candidate if Path(candidate).exists() else None)
        if exe:
            subprocess.run([exe, '--headless', '--convert-to', 'pdf', '--outdir', str(out_dir), str(docx_path)],
                           check=True, capture_output=True)
            return Path(out_dir) / (Path(docx_path).stem + '.pdf')
    raise SystemExit('LibreOffice (soffice) is required to produce the PDF')


def license_holder():
    """The document author is the copyright holder named in this repository's LICENSE."""
    text = (Path(__file__).resolve().parents[1] / 'LICENSE').read_text(encoding='utf-8')
    match = re.search(r'Copyright \(c\) \d{4} (.+)', text)
    return match.group(1).strip() if match else ''


def clean_metadata(doc, meta):
    props = doc.core_properties
    props.author = license_holder()
    props.last_modified_by = license_holder()
    props.title = meta.get('title', '')
    props.subject = ''
    props.keywords = ''
    props.comments = ''
    props.category = ''


def strip_application(docx_path):
    """Removes the application name that the default template carries in docProps/app.xml."""
    path = Path(docx_path)
    with zipfile.ZipFile(path) as source:
        items = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as target:
        for info, data in items:
            if info.filename == 'docProps/app.xml':
                data = re.sub(rb'<Application>[^<]*</Application>|<AppVersion>[^<]*</AppVersion>', b'', data)
            target.writestr(info, data)


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('out_dir')
    parser.add_argument('--pdf', action='store_true')
    parser.add_argument('--label', help='source path as shown in the footer (defaults to the given path)')
    args = parser.parse_args()
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    source = Path(args.source)
    docx_path = out_dir / (source.stem + '.docx')
    render_docx(source, docx_path, args.label or source.as_posix())
    print(f'written {docx_path}')
    if args.pdf:
        print(f'written {to_pdf(docx_path, out_dir)}')


if __name__ == '__main__':
    sys.exit(main())
