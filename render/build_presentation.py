"""Builds the presentation decks (PowerPoint, 16:9) from their Markdown slide sources.

Source format, one deck per language (orderflow-engine-docs/delivery/presentation.md and presentation.es.md):
  front matter              id, title, status, lang, pair, last-reviewed
  # Title                   the title slide; the first paragraph after it is the subtitle
  ## Heading                one slide per heading, titled with it
  - item / 1. step          the slide's text; about six items keep a slide readable
  ![Caption](image.png)     at most one image per slide, its path relative to the .md file
  > **Notes:** ...          the slide's speaker notes ("> **Notas:** ..." in Spanish), as its last block
Paragraphs, ### headings, notices ("> **Note:** ..."), pipe tables and fenced code are rendered too. The question
line under the title ("> **Question this document answers:** ...") is document metadata and is not shown.

Template: Arial, dark text on white, one accent color (the web client's), an accent line under every title, and
"Orderflow Engine" with the slide number in the footer. A slide with text and an image puts the text on the left and
the image on the right; a slide with only an image shows it full width; a slide with only a title is a section
divider. Images keep their aspect ratio. The build stops with a clear message when an image is missing or the source
breaks the format, before any deck is written.

The output is reproducible: the author is the copyright holder named in LICENSE, the title comes from the front
matter, the created and modified dates are its last-reviewed date, no generator name is recorded and every archive
entry carries that same date, so the same source always gives the same bytes.

Usage (from the workspace root):
  python orderflow-engine-tools/render/build_presentation.py
  python orderflow-engine-tools/render/build_presentation.py <source.md> <output.pptx>
"""
import argparse
import io
import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from render_document import front_matter, license_holder  # noqa: E402

WORKSPACE = Path(__file__).resolve().parents[2]

# source -> deck, both workspace-relative
DECKS = [
    ('orderflow-engine-docs/delivery/presentation.md',
     'orderflow-engine-docs/deliverables/delivery/presentation.pptx'),
    ('orderflow-engine-docs/delivery/presentation.es.md',
     'orderflow-engine-docs/deliverables/delivery/presentation.es.pptx'),
]

FONT = 'Arial'
FOOTER_TEXT = 'Orderflow Engine'
# The web client's palette (orderflow-engine-web/src/styles.css): slate ink on white, orange as the only accent.
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)
SUBTLE = RGBColor(0x64, 0x74, 0x8B)
RULE = RGBColor(0xE2, 0xE8, 0xF0)
SUNKEN = RGBColor(0xF1, 0xF5, 0xF9)
ACCENT = RGBColor(0xC2, 0x41, 0x0C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# Proofing language of every run, so PowerPoint checks the Spanish deck as Spanish (a Latin American locale).
LANG_TAGS = {'en': 'en-US', 'es': 'es-CO'}

# Geometry: 16:9 at 13.333 x 7.5 in. The default template is 4:3, 10 in wide.
SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)
TEMPLATE_W = 9144000
MARGIN = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN
TITLE_TOP, TITLE_H = Inches(0.35), Inches(0.9)
RULE_TOP = Inches(1.32)
BODY_TOP, BODY_BOTTOM = Inches(1.65), Inches(6.65)
FOOTER_LINE = Inches(6.85)
FOOTER_TOP, FOOTER_H = Inches(6.93), Inches(0.36)
COLUMN_GAP = Inches(0.4)
IMAGE_COLUMN = (Inches(3.2), Inches(6.4))  # the image column follows the image's aspect ratio within these limits
BLOCK_GAP = Inches(0.18)
CAPTION_GAP = Inches(0.08)
CELL_PAD_X, CELL_PAD_Y = Inches(0.08), Inches(0.04)
CODE_PAD = Inches(0.12)
NEST_INDENT = Inches(0.4)
SLIDE_NUMBER_FIELD = '{8A3C2F61-5D4B-4E2A-9C7F-1B2D3E4F5A60}'

TITLE_SIZES = (30, 28, 26, 24, 22)
FULL_SIZES = (24, 22, 20, 18, 16, 14)
COLUMN_SIZES = (20, 18, 16, 14)
CAPTION_SIZE, FOOTER_SIZE, NOTES_SIZE = 12, 11, 12

# Text measurement, in ems: the average Arial advance per character, kept above the 0.45 that slide text measures,
# and a line height above Arial's single spacing (1.15). Fixed numbers, so the layout never depends on host fonts.
EM_REGULAR, EM_BOLD, LINE = 0.5, 0.55, 1.2
EMU_PER_PT = 12700

IMAGE_TYPES = {'.png', '.jpg', '.jpeg', '.gif', '.bmp'}
LIST_ITEM = re.compile(r'^(\s*)([-*]|(\d+)\.) (.*)$')
IMAGE = re.compile(r'^!\[([^\]]*)\]\(([^)\s]+)\)\s*$')
LABEL = re.compile(r'^\*\*([^*]+?)\s*:?\s*\*\*\s*:?\s*(.*)$')
HEADING = re.compile(r'^(#{1,6}) (.*)$')
BLOCK_START = re.compile(r'^(#{1,6} |\||>|```|!\[|\s*([-*]|\d+\.) )')
TABLE_RULE = re.compile(r'^\|[-| :]+\|\s*$')
NOTES_LABELS = {'notes', 'notas'}
QUESTION_LABELS = {'question this document answers', 'pregunta que responde este documento'}


@dataclass
class Para:
    kind: str  # bullet, step, text or heading
    text: str
    level: int = 0
    number: int = 1


@dataclass
class SlideSpec:
    title: str
    line: int
    blocks: list = field(default_factory=list)  # ('text', [Para]), ('table', rows) or ('code', lines)
    image: tuple = None  # (caption, path, pixel size)
    notes: list = field(default_factory=list)


@dataclass
class Deck:
    meta: dict
    label: str
    title: str = ''
    subtitle: str = ''
    question: str = ''
    extra: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    slides: list = field(default_factory=list)


# --- Source --------------------------------------------------------------------------------------------------------

def parse(source, label):
    """Reads the slide source into a Deck. Every format error stops the build with the file and line."""
    raw = io.open(source, encoding='utf-8').read().lstrip('﻿').replace('\r\n', '\n')
    meta, body = front_matter(raw)
    offset = raw.count('\n', 0, len(raw) - len(body))

    def fail(index, message):
        raise SystemExit(f'{label}:{index + offset + 1}: {message}')

    for key in ('title', 'lang', 'last-reviewed'):
        if not meta.get(key):
            raise SystemExit(f'{label}: the front matter needs "{key}"')
    if meta['lang'] not in LANG_TAGS:
        raise SystemExit(f'{label}: lang must be one of {", ".join(LANG_TAGS)}, found "{meta["lang"]}"')
    try:
        datetime.strptime(meta['last-reviewed'], '%Y-%m-%d')
    except ValueError:
        raise SystemExit(f'{label}: last-reviewed must be a YYYY-MM-DD date, found "{meta["last-reviewed"]}"')

    deck = Deck(meta=meta, label=label)
    slide = None
    lines = body.split('\n')

    def need_slide(index, what):
        if slide is None:
            fail(index, f'{what} before the first "## " slide heading; the title slide holds only the title, '
                        f'a subtitle paragraph and notes')

    def add_para(para):
        if slide.blocks and slide.blocks[-1][0] == 'text':
            slide.blocks[-1][1].append(para)
        else:
            slide.blocks.append(('text', [para]))

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        heading = HEADING.match(line)
        if heading:
            level, text = len(heading.group(1)), heading.group(2).strip()
            if level == 1:
                if deck.title:
                    fail(i, 'a second "# " title; a deck has one title slide')
                deck.title = text
            elif not deck.title:
                fail(i, 'the deck must start with a "# " title')
            elif level == 2:
                slide = SlideSpec(title=text, line=i + offset + 1)
                deck.slides.append(slide)
            elif level == 3:
                need_slide(i, 'a "### " heading')
                add_para(Para('heading', text))
            else:
                fail(i, 'unsupported heading level: use "# " for the deck, "## " for slides and "### " inside one')
            i += 1
            continue
        if not deck.title:
            fail(i, 'the deck must start with a "# " title')
        if line.startswith('```'):
            end = i + 1
            while end < len(lines) and not lines[end].startswith('```'):
                end += 1
            if end == len(lines):
                fail(i, 'this code fence is never closed')
            need_slide(i, 'a code block')
            slide.blocks.append(('code', lines[i + 1:end]))
            i = end + 1
            continue
        if line.startswith('|'):
            start, rows = i, []
            while i < len(lines) and lines[i].startswith('|'):
                if not TABLE_RULE.match(lines[i]):
                    rows.append([cell.strip() for cell in lines[i].strip().strip('|').split('|')])
                i += 1
            need_slide(start, 'a table')
            if any(len(row) != len(rows[0]) for row in rows):
                fail(start, 'every row of this table needs the same number of cells')
            slide.blocks.append(('table', rows))
            continue
        image = IMAGE.match(stripped)
        if image:
            need_slide(i, 'an image')
            if slide.image:
                fail(i, f'slide "{slide.title}" already has an image; a slide shows one image')
            slide.image = load_image(image.group(1).strip(), image.group(2), source, label, lambda m: fail(i, m))
            i += 1
            continue
        if line.startswith('>'):
            start, quote = i, []
            while i < len(lines) and lines[i].startswith('>'):
                quote.append(lines[i][1:].strip())
                i += 1
            paragraphs = [' '.join(part.split('\n')).strip() for part in '\n'.join(quote).split('\n\n')]
            paragraphs = [p for p in paragraphs if p]
            if not paragraphs:
                continue
            label_match = LABEL.match(paragraphs[0])
            name = label_match.group(1).strip().lower() if label_match else ''
            if name in NOTES_LABELS:
                notes = [label_match.group(2).strip()] + paragraphs[1:]
                (slide.notes if slide else deck.notes).extend(n for n in notes if n)
            elif name in QUESTION_LABELS and slide is None:
                deck.question = label_match.group(2).strip()
            else:
                need_slide(start, 'a notice')
                for paragraph in paragraphs:
                    add_para(Para('text', paragraph))
            continue
        item = LIST_ITEM.match(line)
        if item:
            need_slide(i, 'a list item')
            level = 1 if len(item.group(1).expandtabs(4)) >= 2 else 0
            text = item.group(4).strip()
            i += 1
            while (i < len(lines) and lines[i].startswith((' ', '\t')) and lines[i].strip()
                   and not LIST_ITEM.match(lines[i])):
                text += ' ' + lines[i].strip()
                i += 1
            kind = 'step' if item.group(3) else 'bullet'
            add_para(Para(kind, text, level, int(item.group(3) or 1)))
            continue
        text = stripped
        i += 1
        while i < len(lines) and lines[i].strip() and not BLOCK_START.match(lines[i]):
            text += ' ' + lines[i].strip()
            i += 1
        if slide is None:
            if deck.subtitle:
                deck.extra.append(text)
            else:
                deck.subtitle = text
        else:
            add_para(Para('text', text))
    if not deck.title:
        raise SystemExit(f'{label}: no "# " title')
    return deck


def load_image(caption, target, source, label, fail):
    path = source.parent / target
    if Path(target).suffix.lower() not in IMAGE_TYPES:
        fail(f'image {target} must be PNG, JPEG, GIF or BMP')
    if not path.is_file():
        fail(f'image not found: {target} (relative to {Path(label).parent.as_posix()}/)')
    try:
        with Image.open(path) as picture:
            size = picture.size
    except OSError as error:
        fail(f'image {target} cannot be read: {error}')
    if not size[0] or not size[1]:
        fail(f'image {target} is empty')
    return caption, path, size


# --- Text measurement ----------------------------------------------------------------------------------------------

def plain(text):
    return re.sub(r'\*\*|`', '', text)


def wrapped_lines(text, width, size, bold=False):
    """Lines a greedy word wrap needs for the text in a box of the given width (EMU) at the given size (pt)."""
    capacity = max(1, int(width / (size * (EM_BOLD if bold else EM_REGULAR) * EMU_PER_PT)))
    lines, used = 1, 0
    for word in plain(text).split():
        if used and used + 1 + len(word) <= capacity:
            used += 1 + len(word)
        elif not used and len(word) <= capacity:
            used = len(word)
        else:
            if used:
                lines += 1
            spans = -(-len(word) // capacity)
            lines += spans - 1
            used = len(word) - (spans - 1) * capacity
    return lines


def line_height(size):
    return int(size * LINE * EMU_PER_PT)


def item_size(para, size):
    return size - 2 if para.level else size


def hanging(para, size):
    """Left margin and hanging indent of a list item, in EMU."""
    if para.kind not in ('bullet', 'step'):
        return 0, 0
    hang = int(Pt(item_size(para, size)) * (1.3 if para.kind == 'step' else 0.85))
    return para.level * NEST_INDENT + hang, hang


def space_before(para, size):
    if para.kind == 'heading':
        return size * 0.7
    return size * (0.25 if para.level else 0.45)


def text_height(paras, width, size):
    total = 0
    for k, para in enumerate(paras):
        s = item_size(para, size)
        margin, _ = hanging(para, size)
        total += wrapped_lines(para.text, width - margin, s, para.kind == 'heading') * line_height(s)
        if k:
            total += int(space_before(para, size) * EMU_PER_PT)
    return total


def small_size(size):
    return max(12, size - 4)


def column_widths(rows, width, size):
    """No column is narrower than its longest word; the rest of the width follows each column's content length."""
    cols = len(rows[0])
    per_char = size * EM_BOLD * EMU_PER_PT
    words = [max((len(w) for row in rows for w in plain(row[c]).split()), default=1) for c in range(cols)]
    longest = [max(len(plain(row[c])) for row in rows) for c in range(cols)]
    minimum = [int(words[c] * per_char) + 2 * CELL_PAD_X for c in range(cols)]
    if sum(minimum) >= width:
        widths = [int(m * width / sum(minimum)) for m in minimum]
    else:
        extra = [max(0, longest[c] - words[c]) for c in range(cols)]
        if not any(extra):
            extra = [1] * cols
        spare = width - sum(minimum)
        widths = [int(minimum[c] + spare * extra[c] / sum(extra)) for c in range(cols)]
    widths[-1] += width - sum(widths)
    return widths


def table_layout(rows, width, size):
    widths = column_widths(rows, width, size)
    heights = [max(wrapped_lines(cell, widths[c] - 2 * CELL_PAD_X, size, r == 0) for c, cell in enumerate(row))
               * line_height(size) + 2 * CELL_PAD_Y for r, row in enumerate(rows)]
    return widths, heights


def code_height(lines, width, size):
    usable = width - 2 * CODE_PAD
    return sum(wrapped_lines(line or ' ', usable, size) for line in lines) * line_height(size) + 2 * CODE_PAD


def block_height(block, width, size):
    kind, content = block
    if kind == 'text':
        return text_height(content, width, size)
    if kind == 'table':
        return sum(table_layout(content, width, small_size(size))[1])
    return code_height(content, width, small_size(size))


# --- Deck ----------------------------------------------------------------------------------------------------------

class DeckWriter:
    def __init__(self, deck):
        self.deck = deck
        self.lang = LANG_TAGS[deck.meta['lang']]
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        # A custom size has no "type": the template's screen4x3 would make PowerPoint report the wrong format.
        self.prs.part._element.find(qn('p:sldSz')).attrib.pop('type', None)
        widen_template(self.prs)

    def warn(self, line, message):
        print(f'{self.deck.label}:{line}: warning: {message}', file=sys.stderr)

    # Runs and shapes

    def style(self, run, size, color=INK, bold=False):
        font = run.font
        font.name = FONT
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = color
        run._r.get_or_add_rPr().set('lang', self.lang)

    def add_runs(self, paragraph, text, size, color=INK, bold=False, code_color=ACCENT):
        """Supports **bold** and `code` (Arial too, in the accent color)."""
        for token in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
            if not token:
                continue
            run = paragraph.add_run()
            if token.startswith('**') and token.endswith('**') and len(token) > 4:
                run.text = token[2:-2]
                self.style(run, size, color, True)
            elif token.startswith('`') and token.endswith('`') and len(token) > 2:
                run.text = token[1:-1]
                self.style(run, size, code_color, bold)
            else:
                run.text = token
                self.style(run, size, color, bold)

    def textbox(self, slide, left, top, width, height, name, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(left, top, width, height)
        box.name = name
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        return box

    def rect(self, slide, left, top, width, height, color, name):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def layout(self, name):
        for layout in self.prs.slide_layouts:
            if layout.name == name:
                return layout
        raise SystemExit(f'the default template has no "{name}" layout')

    def place_title(self, shape, text, left, top, width, height, sizes, anchor, line):
        shape.left, shape.top, shape.width, shape.height = left, top, width, height
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        size = next((s for s in sizes if wrapped_lines(text, width, s, True) * line_height(s) <= height), None)
        if size is None:
            size = sizes[-1]
            self.warn(line, f'title "{text}" does not fit at {size} pt; shorten it')
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        self.add_runs(paragraph, text, size, INK, True)

    def accent_rule(self, slide, top):
        self.rect(slide, MARGIN, top, CONTENT_W, Pt(1), RULE, 'Title rule')
        self.rect(slide, MARGIN, top - Pt(1.5), Inches(1.2), Pt(4), ACCENT, 'Accent line')

    def footer(self, slide, number):
        self.rect(slide, MARGIN, FOOTER_LINE, CONTENT_W, Pt(0.75), RULE, 'Footer line')
        name = self.textbox(slide, MARGIN, FOOTER_TOP, Inches(6), FOOTER_H, 'Footer', MSO_ANCHOR.MIDDLE)
        self.add_runs(name.text_frame.paragraphs[0], FOOTER_TEXT, FOOTER_SIZE, SUBTLE)
        box = self.textbox(slide, SLIDE_W - MARGIN - Inches(2), FOOTER_TOP, Inches(2), FOOTER_H, 'Slide number',
                           MSO_ANCHOR.MIDDLE)
        paragraph = box.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        # A slide-number field, so the number follows the slide if the deck is reordered; the text is its fallback.
        run = paragraph.add_run()
        self.style(run, FOOTER_SIZE, SUBTLE)
        field_element = OxmlElement('a:fld')
        field_element.set('id', SLIDE_NUMBER_FIELD)
        field_element.set('type', 'slidenum')
        field_element.append(run._r.find(qn('a:rPr')))
        text = OxmlElement('a:t')
        text.text = str(number)
        field_element.append(text)
        paragraph._p.remove(run._r)
        paragraph._p.append(field_element)

    def notes(self, slide, notes):
        if not notes:
            return
        frame = slide.notes_slide.notes_text_frame
        for k, text in enumerate(notes):
            paragraph = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
            self.add_runs(paragraph, text, NOTES_SIZE, INK)

    # Slides

    def title_slide(self):
        deck = self.deck
        slide = self.prs.slides.add_slide(self.layout('Title Slide'))
        self.rect(slide, 0, 0, Inches(0.28), SLIDE_H, ACCENT, 'Accent band')
        left = Inches(1.1)
        width = SLIDE_W - left - MARGIN
        self.place_title(slide.shapes.title, deck.title, left, Inches(1.6), width, Inches(2.0), (44, 40, 36, 32, 28),
                         MSO_ANCHOR.BOTTOM, 1)
        self.rect(slide, left, Inches(3.85), Inches(1.6), Pt(5), ACCENT, 'Accent line')
        subtitle = next(ph for ph in slide.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE)
        text = deck.subtitle or deck.question
        if text:
            subtitle.left, subtitle.top, subtitle.width, subtitle.height = left, Inches(4.1), width, Inches(1.2)
            frame = subtitle.text_frame
            frame.word_wrap = True
            frame.auto_size = MSO_AUTO_SIZE.NONE
            frame.vertical_anchor = MSO_ANCHOR.TOP
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT
            self.add_runs(paragraph, text, 22, MUTED)
        else:
            subtitle._element.getparent().remove(subtitle._element)
        # The byline names the author and the review date, unless the source already names the author.
        lines = list(deck.extra)
        author = license_holder()
        if author and not any(author in line for line in [deck.subtitle] + lines):
            lines.append(f'{author} · {deck.meta["last-reviewed"]}')
        if lines:
            frame = self.textbox(slide, left, Inches(5.45), width, Inches(1.0), 'Byline').text_frame
            for k, line in enumerate(lines):
                paragraph = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
                self.add_runs(paragraph, line, 16, SUBTLE)
        self.notes(slide, deck.notes)

    def content_slide(self, spec, number):
        slide = self.prs.slides.add_slide(self.layout('Title Only'))
        if not spec.blocks and not spec.image:
            # A title alone is a section divider.
            self.place_title(slide.shapes.title, spec.title, MARGIN, Inches(2.4), CONTENT_W, Inches(1.5),
                             (40, 36, 32, 28), MSO_ANCHOR.BOTTOM, spec.line)
            self.accent_rule(slide, Inches(4.1))
        else:
            self.place_title(slide.shapes.title, spec.title, MARGIN, TITLE_TOP, CONTENT_W, TITLE_H, TITLE_SIZES,
                             MSO_ANCHOR.BOTTOM, spec.line)
            self.accent_rule(slide, RULE_TOP)
            height = BODY_BOTTOM - BODY_TOP
            if spec.image and spec.blocks:
                _, _, (pixels_w, pixels_h) = spec.image
                room = height - self.caption_height(spec.image[0], IMAGE_COLUMN[1])
                column = int(min(max(room * pixels_w / pixels_h, IMAGE_COLUMN[0]), IMAGE_COLUMN[1]))
                text_width = CONTENT_W - COLUMN_GAP - column
                self.place_blocks(slide, spec, MARGIN, BODY_TOP, text_width, height, COLUMN_SIZES)
                self.place_image(slide, spec.image, MARGIN + text_width + COLUMN_GAP, BODY_TOP, column, height)
            elif spec.image:
                self.place_image(slide, spec.image, MARGIN, BODY_TOP, CONTENT_W, height)
            else:
                self.place_blocks(slide, spec, MARGIN, BODY_TOP, CONTENT_W, height, FULL_SIZES)
        self.footer(slide, number)
        self.notes(slide, spec.notes)

    def place_blocks(self, slide, spec, left, top, width, height, sizes):
        chosen = None
        for size in sizes:
            heights = [block_height(block, width, size) for block in spec.blocks]
            if sum(heights) + BLOCK_GAP * (len(heights) - 1) <= height:
                chosen = size
                break
        if chosen is None:
            chosen = sizes[-1]
            self.warn(spec.line, f'slide "{spec.title}" does not fit at {chosen} pt; split it or shorten its text')
        y = top
        for block in spec.blocks:
            block_h = block_height(block, width, chosen)
            kind, content = block
            if kind == 'text':
                self.text_block(slide, content, left, y, width, block_h, chosen)
            elif kind == 'table':
                self.table_block(slide, content, left, y, width, small_size(chosen))
            else:
                self.code_block(slide, content, left, y, width, block_h, small_size(chosen))
            y += block_h + BLOCK_GAP

    def text_block(self, slide, paras, left, top, width, height, size):
        frame = self.textbox(slide, left, top, width, height, 'Text').text_frame
        for k, para in enumerate(paras):
            paragraph = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
            if k:
                paragraph.space_before = Pt(space_before(para, size))
            if para.kind in ('bullet', 'step'):
                set_bullet(paragraph, para, size)
            self.add_runs(paragraph, para.text, item_size(para, size), INK, para.kind == 'heading')

    def table_block(self, slide, rows, left, top, width, size):
        widths, heights = table_layout(rows, width, size)
        shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, sum(heights))
        shape.name = 'Table'
        table = shape.table
        table.first_row = True
        table.horz_banding = False
        for c, column_width in enumerate(widths):
            table.columns[c].width = column_width
        for r, row in enumerate(rows):
            table.rows[r].height = heights[r]
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = ACCENT if r == 0 else (WHITE if r % 2 else SUNKEN)
                cell.margin_left = cell.margin_right = CELL_PAD_X
                cell.margin_top = cell.margin_bottom = CELL_PAD_Y
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.text_frame.word_wrap = True
                paragraph = cell.text_frame.paragraphs[0]
                if r == 0:
                    self.add_runs(paragraph, plain(value), size, WHITE, True)
                else:
                    self.add_runs(paragraph, value, size, INK)

    def code_block(self, slide, lines, left, top, width, height, size):
        box = self.textbox(slide, left, top, width, height, 'Code')
        box.fill.solid()
        box.fill.fore_color.rgb = SUNKEN
        frame = box.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = CODE_PAD
        for k, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = line or ' '
            self.style(run, size, INK)

    def caption_height(self, caption, width):
        if not caption:
            return 0
        return wrapped_lines(caption, width, CAPTION_SIZE) * line_height(CAPTION_SIZE) + CAPTION_GAP

    def place_image(self, slide, image, left, top, width, height):
        caption, path, (pixels_w, pixels_h) = image
        caption_h = self.caption_height(caption, width)
        scale = min(width / pixels_w, (height - caption_h) / pixels_h)
        image_w, image_h = int(pixels_w * scale), int(pixels_h * scale)
        y = top + (height - image_h - caption_h) // 2
        picture = slide.shapes.add_picture(str(path), left + (width - image_w) // 2, y, image_w, image_h)
        picture.name = 'Image'
        picture._element.nvPicPr.cNvPr.set('descr', plain(caption) or path.stem)
        picture.line.color.rgb = RULE
        picture.line.width = Pt(0.75)
        if caption:
            box = self.textbox(slide, left, y + image_h + CAPTION_GAP, width, caption_h - CAPTION_GAP, 'Caption')
            paragraph = box.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            self.add_runs(paragraph, caption, CAPTION_SIZE, MUTED)

    def save(self, output):
        self.title_slide()
        for number, spec in enumerate(self.deck.slides, start=2):
            self.content_slide(spec, number)
        date = datetime.strptime(self.deck.meta['last-reviewed'], '%Y-%m-%d')
        author = license_holder()
        props = self.prs.core_properties
        props.author = author
        props.last_modified_by = author
        props.title = self.deck.meta['title']
        props.subject = props.keywords = props.comments = props.category = ''
        props.language = self.lang
        props.revision = 1
        props.created = props.modified = date
        output.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(output))
        notes = sum(1 for slide in self.prs.slides if slide.has_notes_slide)
        finalize(output, date, len(self.prs.slides), notes)
        return len(self.prs.slides)


def set_bullet(paragraph, para, size):
    """A bullet (accent dot or dash) or a step number (from the source) with a hanging indent."""
    margin, hang = hanging(para, size)
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set('marL', str(margin))
    p_pr.set('indent', str(-hang))
    color = OxmlElement('a:buClr')
    srgb = OxmlElement('a:srgbClr')
    srgb.set('val', str(ACCENT))
    color.append(srgb)
    percent = OxmlElement('a:buSzPct')
    percent.set('val', '100000')
    font = OxmlElement('a:buFont')
    font.set('typeface', FONT)
    if para.kind == 'step':
        marker = OxmlElement('a:buAutoNum')
        marker.set('type', 'arabicPeriod')
        if para.number != 1:
            marker.set('startAt', str(para.number))
    else:
        marker = OxmlElement('a:buChar')
        marker.set('char', '•' if para.level == 0 else '–')
    for element in (color, percent, font, marker):
        p_pr.append(element)


def widen_template(prs):
    """Stretches the 4:3 master and layouts to 16:9, so a slide added later in PowerPoint fits the new width."""
    ratio = SLIDE_W / TEMPLATE_W
    for owner in [prs.slide_master, *prs.slide_layouts]:
        for shape in owner.shapes:
            if getattr(shape._element, 'xfrm', None) is None:
                continue  # inherits its position from the master, which is stretched already
            shape.left = int(shape.left * ratio)
            shape.width = int(shape.width * ratio)


APP_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           '<TotalTime>0</TotalTime><PresentationFormat>Widescreen</PresentationFormat>'
           '<Slides>{slides}</Slides><Notes>{notes}</Notes><HiddenSlides>0</HiddenSlides>'
           '<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc>'
           '<HyperlinksChanged>false</HyperlinksChanged></Properties>')


def finalize(path, date, slides, notes):
    """Rewrites the package so it is reproducible and carries no template leftovers.

    The extended properties are replaced (the template names an application, its version and a 4:3 format), the
    template's blank 4:3 thumbnail is dropped, the theme fonts become Arial and its first accent the deck's accent,
    and every archive entry gets the document date instead of the build time.
    """
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    stamp = (date.year, date.month, date.day, 0, 0, 0)
    with zipfile.ZipFile(path, 'w') as target:
        for name, data in entries:
            if name == 'docProps/thumbnail.jpeg':
                continue
            if name == '_rels/.rels':
                data = re.sub(rb'<Relationship [^>]*metadata/thumbnail"[^>]*/>', b'', data)
            elif name == 'docProps/app.xml':
                data = APP_XML.format(slides=slides, notes=notes).encode('utf-8')
            elif name.startswith('ppt/theme/'):
                data = re.sub(rb'(<a:(?:major|minor)Font><a:latin typeface=")[^"]*', rb'\1' + FONT.encode(), data)
                data = re.sub(rb'(<a:accent1><a:srgbClr val=")[0-9A-Fa-f]{6}', rb'\1' + str(ACCENT).encode(), data)
            info = zipfile.ZipInfo(name, date_time=stamp)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            target.writestr(info, data)


def main():
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors='replace')
    parser = argparse.ArgumentParser(description='Builds the PowerPoint decks from their Markdown slide sources.')
    parser.add_argument('source', nargs='?', help='slide source (.md); builds both decks when omitted')
    parser.add_argument('output', nargs='?', help='deck to write (.pptx)')
    args = parser.parse_args()
    if bool(args.source) != bool(args.output):
        parser.error('give both <source.md> and <output.pptx>, or neither')
    if args.source:
        jobs = [(Path(args.source), Path(args.output), Path(args.source).as_posix(), args.output)]
    else:
        jobs = [(WORKSPACE / source, WORKSPACE / output, source, output) for source, output in DECKS]
    decks = []
    for source, output, label, shown in jobs:
        if not source.is_file():
            raise SystemExit(f'{label}: source not found')
        decks.append((parse(source, label), output, shown))  # every source is checked before any deck is written
    for deck, output, shown in decks:
        count = DeckWriter(deck).save(output)
        print(f'written {shown} ({count} slides)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
